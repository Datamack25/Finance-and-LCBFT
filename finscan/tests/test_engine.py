"""Tests rapides : python -m pytest -q

Les tests d'intégration lisent les PDF placés dans tests/fixtures/
(non versionnés : documents clients). Ils sont ignorés s'ils sont absents.
"""
from pathlib import Path

import pytest

from engine.ingest import Line, Word
from engine.kyc import luhn_ok, same_person
from engine.numbers import parse_line

FIX = Path(__file__).parent / "fixtures"


def _line(*tokens):
    """tokens : (texte, x0) ; largeur ≈ 7 px par caractère, hauteur 10 px."""
    return Line(1, [Word(t, x, x + 7 * len(t), 10) for t, x in tokens])


def test_french_thousands_and_columns():
    pl = parse_line(_line(("Immobilisations", 0), ("corporelles", 110), ("551", 400), ("361", 428), ("132", 520), ("501", 548)))
    assert pl.label == "Immobilisations corporelles"
    assert [n.value for n in pl.nums] == [551361, 132501]


def test_negative_and_us_format():
    pl = parse_line(_line(("Résultat", 0), ("-", 390), ("184", 400), ("486", 428)))
    assert pl.nums[0].value == -184486
    pl = parse_line(_line(("Other", 0), ("(527)", 400)))
    assert pl.nums[0].value == -527
    pl = parse_line(_line(("Net", 0), ("sales", 30), ("$", 380), ("254,940", 390)))
    assert pl.nums[0].value == 254940


def test_note_reference_ignored():
    pl = parse_line(_line(("Autres", 0), ("actifs", 50), ("2.05", 300), ("12", 400), ("281", 420)))
    assert [n.value for n in pl.nums if not n.note] == [12281]
    assert pl.nums[0].note  # gardé mais marqué : écarté s'il ne tombe sous aucune colonne


def test_siren_luhn():
    assert luhn_ok("393548987")
    assert not luhn_ok("593548987")


def test_same_person_spelling_variant():
    ok, variant = same_person("Madame Michèle NAON", "Madame Michelle AOUIZERAT épouse NAON")
    assert ok and variant
    ok, _ = same_person("Mme Johanna RIAHI", "Madame Johanna NAON épouse RIAHI")
    assert ok
    ok, _ = same_person("Paul NAON", "Laetitia NAON")
    assert not ok


@pytest.mark.skipif(not list(FIX.glob("*.pdf")), reason="aucun PDF dans tests/fixtures")
def test_pdfs_balance():
    from engine.dossier import build_dossier
    for pdf in FIX.glob("*.pdf"):
        d = build_dossier(pdf.read_bytes(), pdf.name)
        res = d.analyse()
        bal = next(c for c in res["controls"] if c["contrôle"].startswith("Bilan"))
        assert bal["statut"] in ("ok", "incomplet"), (pdf.name, bal)


# ---------------------------------------------------------------------------
# Multi-exercices : PDF fictifs générés à la volée
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def demo(tmp_path_factory):
    from tests.fixtures_gen import make_all
    from engine.company import company_key, merge, new_company
    from engine.dossier import build_dossier
    out = make_all(tmp_path_factory.mktemp("demo"))
    docs, comps = {}, {}
    for paths in out.values():
        for p in paths:
            d = build_dossier(p.read_bytes(), p.name)
            docs[d.id] = d
            k = company_key(d)
            if k in comps:
                comps[k].doc_ids.append(d.id)
            else:
                comps[k] = new_company(d)
    return {c.name.split()[0].upper(): merge(c, docs).analyse() for c in comps.values()}


def test_pcg_three_documents_merge_into_four_years(demo):
    f = demo["NOVALIS"]["fin"]
    assert f.periods == ["2024", "2023", "2022", "2021"]
    assert f.get("revenue", "2021") == 1_800_000
    assert f.get("revenue", "2024") == 2_528_870
    tcam = next(r for r in demo["NOVALIS"]["ratios"] if r["ratio"].startswith("Chiffre d'affaires — TCAM"))
    assert abs(tcam["N"] - 0.12) < 1e-3
    bal = next(c for c in demo["NOVALIS"]["controls"] if c["contrôle"].startswith("Bilan"))
    assert bal["statut"] == "ok"


def test_ifrs_french_three_columns(demo):
    f = demo["GROUPE"]["fin"]
    assert f.periods == ["2024", "2023", "2022"]
    assert f.get("revenue", "2022") == 1_250e6
    assert f.get("equity", "2024") == 1_230e6
    assert f.get("minority", "2024") == 25e6


def test_ifrs_english_and_tax_sign(demo):
    f = demo["NORTHWIND"]["fin"]
    assert f.get("net_income", "2023") == 74e6
    assert f.get("equity", "2022") == 308e6
    ctl = next(c for c in demo["NORTHWIND"]["controls"] if c["contrôle"].startswith("Résultat net"))
    assert ctl["statut"] == "ok"


def test_us_10k_three_years(demo):
    f = demo["CONTOSO"]["fin"]
    assert [f.get("net_income", p) for p in ("2022", "2023", "2024")] == [1169e6, 1481e6, 1746e6]
    assert f.get("total_assets", "2024") == 9630e6
    assert f.get("eps_diluted", "2024") == 3.49


def test_exports_word_and_excel(demo, tmp_path):
    """Les exports se construisent ; le classeur contient les onglets attendus."""
    import io
    from openpyxl import load_workbook
    from engine.company import company_key, merge, new_company
    from engine.dossier import build_dossier
    from engine.export_docx import build_docx
    from engine.export_xlsx import build_xlsx
    from tests.fixtures_gen import make_novalis
    docs = {}
    for p in make_novalis(tmp_path):
        d = build_dossier(p.read_bytes(), p.name)
        docs[d.id] = d
    c = new_company(next(iter(docs.values())))
    c.doc_ids = list(docs)
    v = merge(c, docs)
    res = v.analyse()
    assert build_docx(v, res)[:2] == b"PK"
    wb = load_workbook(io.BytesIO(build_xlsx(v, res, docs)))
    for name in ("Synthèse", "Bilan", "Compte de résultat", "Flux de trésorerie", "Ratios", "Liasses et documents",
                 "Actionnaires", "Graphiques actionnaires", "LCB-FT"):
        assert name in wb.sheetnames
    assert wb["Actionnaires"]["A5"].value == "Karim BENALI"
