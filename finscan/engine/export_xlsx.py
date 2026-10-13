"""Classeur Excel complet d'une société.

Onglets : Synthèse · Bilan · Compte de résultat · Flux de trésorerie (publié ou
reconstitué) · Liasses et documents (toutes les lignes lues, document par
document) · Ratios · Actionnaires · Graphiques actionnaires · LCB-FT.

Les montants lus sont des valeurs ; les variations, TCAM, agrégats et ratios
sont des formules Excel, recalculées si l'on corrige un chiffre.
"""
from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .fields import FORMAT_NAMES, LABELS
from .kyc import is_legal_entity, same_person
from .report import EXTRA_LABELS, label

NAVY = "1F3A5F"
SOFT = "E8EDF4"
F_BASE = Font(name="Arial", size=10)
F_BOLD = Font(name="Arial", size=10, bold=True)
F_HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name="Arial", size=14, bold=True, color=NAVY)
F_NOTE = Font(name="Arial", size=9, italic=True, color="5A6475")
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_SOFT = PatternFill("solid", fgColor=SOFT)
FILL_INPUT = PatternFill("solid", fgColor="FFF8E1")
THIN = Side(style="thin", color="D9DEE6")
BORDER = Border(bottom=THIN)
NUM = '#,##0;[Red]-#,##0;"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
RATIO_X = '0.00"x";[Red]-0.00"x"'

BILAN_ACTIF = ["intangible", "tangible", "financial_assets", "lt_securities", "fixed_assets", "inventories", "receivables",
               "other_receivables", "interbank_assets", "other_assets", "marketable_securities", "cash", "prepaid",
               "current_assets", "total_assets"]
BILAN_PASSIF = ["share_capital", "share_premium", "reserves", "retained_carry", "net_income_bs", "regulated_prov", "minority",
                "equity", "provisions", "financial_debt", "commercial_paper", "interbank_debt", "client_funds",
                "customer_advances", "payables", "tax_social_debts", "other_liabilities", "deferred_income",
                "current_liabilities", "total_liabilities", "total_liabilities_equity"]
RESULTAT = ["revenue_direct", "revenue_parts", "revenue", "operating_revenue_total", "interest_income", "interest_expense",
            "fee_income", "fee_expense", "pnb", "cogs", "gross_profit", "purchases", "external_charges", "taxes", "wages",
            "social_charges", "personnel", "rd", "sga", "other_opex", "d_and_a", "provisions_charge",
            "operating_expenses_total", "rbe", "cost_of_risk", "ebit", "financial_income", "financial_expense", "other_income",
            "exceptional_income", "exceptional_expense", "pretax", "income_tax", "net_income", "eps_basic", "eps_diluted",
            "shares_diluted"]
FLUX = ["cfo", "sbc", "capex", "cfi", "dividends_paid", "buybacks", "debt_repaid", "cff", "taxes_paid"]
ANNEXE = ["safeguarded", "acquisition_accounts", "reserve_accounts"]
# agrégats écrits en formules (composantes, signe)
AGGREGATES = {
    "liquid": [("cash", 1), ("marketable_securities", 1)],
    "financial_debt_total": [("financial_debt", 1), ("commercial_paper", 1), ("interbank_debt", 1)],
    "ebitda": [("ebit", 1), ("d_and_a", 1), ("provisions_charge", 1)],
    "general_expenses": [("external_charges", 1), ("personnel", 1), ("taxes", 1), ("other_opex", 1)],
}


def _sheet_ref(ws) -> str:
    return f"'{ws.title}'" if re.search(r"[^A-Za-z0-9_]", ws.title) else ws.title


class Book:
    def __init__(self, view, res):
        self.v, self.res = view, res
        self.f = res["fin"]
        self.periods = list(reversed(self.f.periods))  # chronologique
        self.wb = Workbook()
        self.refs: dict[str, dict[str, str]] = {}  # poste -> période -> référence de cellule
        self.cur = view.ex.currency

    # ------------------------------------------------------------------ outils
    def _title(self, ws, text, sub=""):
        ws["A1"] = text
        ws["A1"].font = F_TITLE
        if sub:
            ws["A2"] = sub
            ws["A2"].font = F_NOTE
        ws.sheet_view.showGridLines = False

    def _header(self, ws, row, values, start_col=1):
        for i, val in enumerate(values):
            c = ws.cell(row=row, column=start_col + i, value=val)
            c.font, c.fill = F_HEAD, FILL_HEAD
            c.alignment = Alignment(horizontal="center" if i else "left", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 30

    def _widths(self, ws, widths: dict[str, float]):
        for col, w in widths.items():
            ws.column_dimensions[col].width = w

    def _cell(self, ws, row, col, value, fmt=None, font=F_BASE, fill=None, comment=None):
        c = ws.cell(row=row, column=col, value=value)
        c.font = font
        if fmt:
            c.number_format = fmt
        if fill:
            c.fill = fill
        if comment:
            c.comment = Comment(comment, "FinScan")
        c.border = BORDER
        return c

    # --------------------------------------------------------- états financiers
    def statement(self, title: str, blocks: list[tuple[str, list[str]]], note: str):
        ws = self.wb.create_sheet(title)
        n = len(self.periods)
        self._title(ws, f"{self.v.name} — {title}", note)
        hdr = ["Poste"] + self.periods + [f"Δ {p}" for p in self.periods[1:]] + (["TCAM"] if self._annual_count() >= 3 else []) + \
              ["Origine", "Source (document, page, libellé)"]
        self._header(ws, 4, hdr)
        row = 5
        sref = _sheet_ref(ws)
        for btitle, keys in blocks:
            present = [k for k in keys if k in self.f.v and any(self.f.get(k, p) is not None for p in self.periods)]
            present += [k for k in AGGREGATES if k in keys and k not in present and k in self.f.v]
            if "revenue" in present and ("revenue_direct" in present or "revenue_parts" in present or "pnb" in present):
                pass  # la ligne « retenu » devient une formule pointant vers la ligne publiée
            if not present:
                continue
            c = ws.cell(row=row, column=1, value=btitle)
            c.font, c.fill = F_BOLD, FILL_SOFT
            for col in range(2, len(hdr) + 1):
                ws.cell(row=row, column=col).fill = FILL_SOFT
            row += 1
            for k in present:
                lab_k = "Chiffre d'affaires retenu (formule)" if k == "revenue" else (EXTRA_LABELS.get(k) or LABELS.get(k, k))
                self._cell(ws, row, 1, lab_k, font=F_BOLD if k.startswith("total") or k in
                           ("equity", "net_income", "revenue", "ebit", "pnb") else F_BASE)
                self.refs.setdefault(k, {})
                for j, p in enumerate(self.periods):
                    col = 2 + j
                    addr = f"{get_column_letter(col)}{row}"
                    agg = AGGREGATES.get(k)
                    if k == "revenue":  # chiffre d'affaires retenu = ligne publiée, sinon somme des ventes, sinon PNB
                        src_k = next((x for x in ("revenue_direct", "revenue_parts", "pnb") if p in self.refs.get(x, {})), None)
                        agg = [(src_k, 1)] if src_k else None
                    parts = [(ck, sg) for ck, sg in agg if p in self.refs.get(ck, {})] if agg else []
                    if agg and parts and self.f.origin.get(k) == "calculé":
                        formula = "=" + "+".join(("" if sg > 0 else "-") + self.refs[ck][p] for ck, sg in parts)
                        self._cell(ws, row, col, formula.replace("+-", "-"), NUM)
                    else:
                        val = self.f.get(k, p)
                        fmt = "0.00" if k.startswith("eps") else NUM
                        self._cell(ws, row, col, val, fmt, fill=FILL_INPUT if self.f.origin.get(k) == "saisi" else None)
                    self.refs[k][p] = f"{sref}!${get_column_letter(col)}${row}"
                # variations (formules)
                for j in range(1, n):
                    col = 2 + n - 1 + j
                    cur_c, prev_c = f"{get_column_letter(2 + j)}{row}", f"{get_column_letter(1 + j)}{row}"
                    if self.f.months.get(self.periods[j], 12) != self.f.months.get(self.periods[j - 1], 12):
                        self._cell(ws, row, col, None)
                        continue
                    self._cell(ws, row, col, f'=IF(OR({prev_c}="",{prev_c}=0,{cur_c}=""),"",({cur_c}-{prev_c})/ABS({prev_c}))', PCT)
                col = 2 + n - 1 + n
                if self._annual_count() >= 3:
                    first, last = self._annual_bounds()
                    fc, lc = f"{get_column_letter(2 + first)}{row}", f"{get_column_letter(2 + last)}{row}"
                    years = int(self.periods[last][:4]) - int(self.periods[first][:4])
                    self._cell(ws, row, col, f'=IF(AND(N({fc})>0,N({lc})>0),({lc}/{fc})^(1/{years})-1,"")', PCT)
                    col += 1
                self._cell(ws, row, col, self.f.origin.get(k, ""))
                src = self.v.ex.values.get(k).source if k in self.v.ex.values else ("formule" if k in AGGREGATES else "calcul")
                self._cell(ws, row, col + 1, src)
                row += 1
            row += 1
        widths = {"A": 42}
        for j in range(len(hdr) - 1):
            widths[get_column_letter(2 + j)] = 16
        widths[get_column_letter(len(hdr) - 1)] = 11
        widths[get_column_letter(len(hdr))] = 60
        self._widths(ws, widths)
        ws.freeze_panes = "B5"
        ws.cell(row=row + 1, column=1, value="Montants en " + self.cur + " (unités). Cellules jaunes : saisies de l'analyste. "
                "Δ = variation sur l'exercice précédent ; TCAM = taux de croissance annuel moyen. Les variations, TCAM et "
                "agrégats sont des formules : corriger un montant met tout à jour.").font = F_NOTE
        return ws

    def _annual_idx(self):
        return [i for i, p in enumerate(self.periods) if self.f.months.get(p, 12) == 12]

    def _annual_count(self):
        return len(self._annual_idx())

    def _annual_bounds(self):
        idx = self._annual_idx()
        return idx[0], idx[-1]

    def flux(self):
        has_cf = any(k in self.f.v and any(self.f.get(k, p) is not None for p in self.periods) for k in ("cfo", "capex"))
        if has_cf:
            return self.statement("Flux de trésorerie", [("Flux publiés", FLUX + ["fcf"])],
                                  "Tableau des flux tel que publié. Free cash-flow = flux d'exploitation − investissements.")
        # Flux reconstitués (méthode indirecte simplifiée) à partir du bilan et du compte de résultat
        ws = self.wb.create_sheet("Flux de trésorerie")
        self._title(ws, f"{self.v.name} — Flux de trésorerie reconstitués",
                    "Aucun tableau de flux dans les documents : estimation par la méthode indirecte à partir du bilan et du "
                    "compte de résultat (formules). À utiliser comme ordre de grandeur.")
        self._header(ws, 4, ["Poste"] + self.periods + ["Formule"])
        r = lambda k, p: self.refs.get(k, {}).get(p)
        lines = [
            ("Résultat net", lambda p, q: r("net_income", p), "Résultat net"),
            ("Dotations aux amortissements et provisions (+)", lambda p, q: "+".join(x for x in (r("d_and_a", p), r("provisions_charge", p)) if x), "DAP"),
            ("Capacité d'autofinancement (CAF)", "sum2", "Résultat net + DAP"),
            ("Variation du BFR (−)", "bfr", "Δ (stocks + clients + autres créances − fournisseurs − dettes fiscales et sociales − autres dettes)"),
            ("Flux d'exploitation estimé", "cfo", "CAF − ΔBFR"),
            ("Investissements estimés (−)", "capex", "Δ immobilisations nettes + dotations aux amortissements"),
            ("Free cash-flow estimé", "fcf", "Flux d'exploitation − investissements"),
            ("Variation de la trésorerie et des placements (bilan)", "dcash", "Δ (disponibilités + VMP)"),
        ]
        bfr_pos = ["inventories", "receivables", "other_receivables"]
        bfr_neg = ["payables", "tax_social_debts", "other_liabilities", "customer_advances"]
        immo = ["intangible", "tangible", "financial_assets"]
        for i, (lab, fn, expl) in enumerate(lines):
            row = 5 + i
            self._cell(ws, row, 1, lab, font=F_BOLD if lab.startswith(("Capacité", "Flux", "Free")) else F_BASE)
            for j, p in enumerate(self.periods):
                col = 2 + j
                q = self.periods[j - 1] if j else None
                cl = get_column_letter(col)
                val = None
                if fn == "sum2":
                    val = f"=N({cl}5)+N({cl}6)"
                elif fn == "bfr" and q:
                    pos = [x for k in bfr_pos for x in (r(k, p), r(k, q)) if x]
                    terms = []
                    for k in bfr_pos:
                        if r(k, p) and r(k, q):
                            terms.append(f"+({r(k, p)}-{r(k, q)})")
                    for k in bfr_neg:
                        if r(k, p) and r(k, q):
                            terms.append(f"-({r(k, p)}-{r(k, q)})")
                    val = "=" + "".join(terms).lstrip("+") if terms and pos else None
                elif fn == "cfo":
                    val = f'=IF({cl}8="","",N({cl}7)-N({cl}8))' if q else None
                elif fn == "capex" and q:
                    terms = [f"+({r(k, p)}-{r(k, q)})" for k in immo if r(k, p) and r(k, q)]
                    if terms:
                        val = "=" + "".join(terms).lstrip("+") + (f"+N({r('d_and_a', p)})" if r("d_and_a", p) else "")
                elif fn == "fcf":
                    val = f'=IF(OR({cl}9="",{cl}10=""),"",{cl}9-{cl}10)' if q else None
                elif fn == "dcash" and q:
                    terms = [f"+(N({r(k, p)})-N({r(k, q)}))" for k in ("cash", "marketable_securities") if r(k, p)]
                    val = ("=" + "".join(terms).lstrip("+")) if terms else None
                elif callable(fn):
                    ref = fn(p, q)
                    val = f"={ref}" if ref else None
                self._cell(ws, row, col, val, NUM)
            self._cell(ws, row, 2 + len(self.periods), expl, font=F_NOTE)
        self._widths(ws, {"A": 52, **{get_column_letter(2 + j): 16 for j in range(len(self.periods))},
                          get_column_letter(2 + len(self.periods)): 70})
        ws.freeze_panes = "B5"
        return ws

    # ------------------------------------------------------------------ ratios
    def ratios(self):
        ws = self.wb.create_sheet("Ratios")
        self._title(ws, f"{self.v.name} — Ratios", "Formules reliées aux onglets des états financiers quand le ratio est un "
                    "rapport simple ; sinon valeur calculée par la plateforme (commentaire sur la cellule).")
        self._header(ws, 4, ["Groupe", "Ratio"] + self.periods + ["Appréciation (dernier exercice)", "Note"])
        r = lambda k, p: self.refs.get(k, {}).get(p)

        def div(a, b, ann=False):
            def fn(p):
                ra, rb = r(a, p), r(b, p)
                if not ra or not rb:
                    return None
                m = self.f.months.get(p, 12)
                mult = f"*12/{m}" if ann and m != 12 else ""
                return f'=IF(N({rb})=0,"",{ra}{mult}/{rb})'
            return fn

        def grow(k):
            def fn(p):
                i = self.periods.index(p)
                if i == 0:
                    return None
                q = self.periods[i - 1]
                if self.f.months.get(q, 12) != self.f.months.get(p, 12) or not r(k, p) or not r(k, q):
                    return None
                return f'=IF(N({r(k, q)})=0,"",({r(k, p)}-{r(k, q)})/ABS({r(k, q)}))'
            return fn

        formulas = {
            "Croissance du chiffre d'affaires": grow("revenue"), "Croissance du résultat net": grow("net_income"),
            "Croissance des capitaux propres": grow("equity"), "Croissance du total de bilan": grow("total_assets"),
            "Croissance du PNB": grow("pnb"), "Croissance du BPA dilué": grow("eps_diluted"),
            "Marge brute": div("gross_profit", "revenue"), "Marge d'exploitation (EBIT / CA)": div("ebit", "revenue"),
            "Marge d'EBITDA / EBE": div("ebitda", "revenue"), "Marge nette": div("net_income", "revenue"),
            "Marge nette (résultat / PNB)": div("net_income", "pnb"),
            "ROE (rentabilité des capitaux propres)": div("net_income", "equity", True), "ROE": div("net_income", "equity", True),
            "ROA (rentabilité de l'actif)": div("net_income", "total_assets", True), "ROA": div("net_income", "total_assets", True),
            "Liquidité générale (actif circulant / passif courant)": div("current_assets", "current_liabilities"),
            "Liquidité immédiate (trésorerie + VMP / passif courant)": div("liquid", "current_liabilities"),
            "Autonomie financière (capitaux propres / total bilan)": div("equity", "total_assets"),
            "Capitaux propres / total bilan": div("equity", "total_assets"),
            "Gearing (dettes financières / capitaux propres)": div("financial_debt_total", "equity"),
            "Dettes bancaires / capitaux propres": div("interbank_debt", "equity"),
            "Liquidités / total bilan": div("liquid", "total_assets"),
            "Fonds cantonnés / fonds de la clientèle": div("safeguarded", "client_funds"),
            "Rotation de l'actif (CA / total actif)": div("revenue", "total_assets", True),
            "Free cash-flow / CA": div("fcf", "revenue"),
            "Conversion du résultat en cash (CFO / résultat net)": div("cfo", "net_income"),
            "Poids des produits financiers (produits financiers / CA)": div("financial_income", "revenue"),
        }
        row = 5
        for rt in self.res["ratios"]:
            vals = rt.get("values", {})
            if all(vals.get(p) is None for p in self.periods):
                continue
            self._cell(ws, row, 1, rt["groupe"])
            self._cell(ws, row, 2, rt["ratio"])
            fmt = PCT if rt["format"] == "%" else RATIO_X if rt["format"] == "x" else '0" j"' if rt["format"] == "j" else "0.00"
            fn = formulas.get(rt["ratio"])
            for j, p in enumerate(self.periods):
                f_ = fn(p) if fn else None
                if f_ and vals.get(p) is not None:
                    self._cell(ws, row, 3 + j, f_, fmt)
                else:
                    self._cell(ws, row, 3 + j, vals.get(p), fmt,
                               comment="Valeur calculée par la plateforme" if vals.get(p) is not None else None)
            self._cell(ws, row, 3 + len(self.periods), rt["appréciation"])
            self._cell(ws, row, 4 + len(self.periods), rt["note"], font=F_NOTE)
            row += 1
        self._widths(ws, {"A": 16, "B": 58, **{get_column_letter(3 + j): 13 for j in range(len(self.periods))},
                          get_column_letter(3 + len(self.periods)): 16, get_column_letter(4 + len(self.periods)): 70})
        ws.freeze_panes = "C5"

    # --------------------------------------------------------- liasses / documents
    def documents(self, docs: dict):
        ws = self.wb.create_sheet("Liasses et documents")
        self._title(ws, f"{self.v.name} — Toutes les lignes lues, document par document",
                    "Montants dans l'unité de chaque document (colonne Unité). « Non rattachée » : ligne chiffrée lue mais non "
                    "affectée à un poste standard — elle figure ici pour que tous les chiffres soient disponibles.")
        all_p = sorted({p for s in self.v.sources for p in docs[s["doc_id"]].ex.periods} if docs else set(), reverse=True)
        all_p = list(reversed(all_p))
        self._header(ws, 4, ["Document", "Format", "Unité", "Page", "Section", "Libellé lu", "Poste rattaché"] + all_p)
        row = 5
        sec_names = {"A": "Actif", "P": "Passif", "R": "Résultat", "F": "Flux", "X": "Annexe"}
        unit_names = {1.0: "unités", 1000.0: "milliers", 1e6: "millions"}
        for s in self.v.sources:
            d = docs.get(s["doc_id"])
            if not d:
                continue
            ex = d.ex
            unit = f"{unit_names.get(ex.unit, ex.unit)} {ex.currency}"
            rows = []
            for k, val in ex.values.items():
                rows.append((val.page or 0, "", val.source, LABELS.get(k, k), val.vals))
            for u in ex.unmatched:
                rows.append((u["page"], sec_names.get(u.get("section"), ""), u["libellé"], "non rattachée", u.get("vals", {})))
            rows.sort(key=lambda t: (t[0], t[3] == "non rattachée"))
            for page, sec, lib, poste, vals in rows:
                self._cell(ws, row, 1, d.file_name)
                self._cell(ws, row, 2, ex.fmt)
                self._cell(ws, row, 3, unit)
                self._cell(ws, row, 4, page or None)
                self._cell(ws, row, 5, sec)
                self._cell(ws, row, 6, lib)
                self._cell(ws, row, 7, poste, font=F_NOTE if poste == "non rattachée" else F_BASE)
                for j, p in enumerate(all_p):
                    x = vals.get(p)
                    self._cell(ws, row, 8 + j, x, "0.00" if x is not None and abs(x) < 100 and x != int(x) else NUM)
                row += 1
        self._widths(ws, {"A": 34, "B": 8, "C": 14, "D": 7, "E": 10, "F": 52, "G": 34,
                          **{get_column_letter(8 + j): 15 for j in range(len(all_p))}})
        ws.freeze_panes = "H5"
        ws.auto_filter.ref = f"A4:{get_column_letter(7 + len(all_p))}{max(row - 1, 5)}"

    # ------------------------------------------------------------- actionnaires
    def shareholders(self):
        k = self.v.kyc
        ws = self.wb.create_sheet("Actionnaires")
        self._title(ws, f"{self.v.name} — Actionnariat, dirigeants, filiales",
                    "Pourcentages et écarts en formules. Un écart entre part du capital et part des dividendes peut révéler un "
                    "démembrement ou un bénéficiaire économique différent du titulaire des parts.")
        self._header(ws, 4, ["Associé / actionnaire", "Parts ou actions", "% du capital", "Dividendes décidés", "% des dividendes",
                             "Écart (points)", "Nature", "Bénéficiaire effectif (> 25 %)"])
        row = 5
        first = row
        assoc = k.associates
        divs = list(k.dividends_by_person)
        for p in assoc:
            dv = next((d.dividend for d in divs if same_person(d.name, p.name)[0]), None)
            self._cell(ws, row, 1, p.name)
            self._cell(ws, row, 2, p.parts, NUM)
            self._cell(ws, row, 4, dv, NUM)
            row += 1
        last = row - 1
        if assoc:
            has_parts = any(p.parts for p in assoc)
            for i, p in enumerate(assoc):
                rr = first + i
                if has_parts:
                    self._cell(ws, rr, 3, f"=IF(SUM($B${first}:$B${last})=0,\"\",B{rr}/SUM($B${first}:$B${last}))", PCT)
                else:
                    self._cell(ws, rr, 3, (p.pct or 0) / 100, PCT)
                self._cell(ws, rr, 5, f"=IF(SUM($D${first}:$D${last})=0,\"\",D{rr}/SUM($D${first}:$D${last}))", PCT)
                self._cell(ws, rr, 6, f'=IF(E{rr}="","",(E{rr}-C{rr})*100)', '0.0;[Red]-0.0')
                legal = is_legal_entity(p.name)
                self._cell(ws, rr, 7, "personne morale" if legal else "personne physique")
                if legal:
                    self._cell(ws, rr, 8, f'=IF(OR(N(C{rr})>0.25,N(E{rr})>0.25),"non : remonter la chaîne de détention","")')
                else:
                    self._cell(ws, rr, 8, f'=IF(OR(N(C{rr})>0.25,N(E{rr})>0.25),"oui (présumé)","")')
            self._cell(ws, row, 1, "Total", font=F_BOLD)
            self._cell(ws, row, 2, f"=SUM(B{first}:B{last})", NUM, F_BOLD)
            self._cell(ws, row, 3, f"=SUM(C{first}:C{last})", PCT, F_BOLD)
            self._cell(ws, row, 4, f"=SUM(D{first}:D{last})", NUM, F_BOLD)
            row += 2
        else:
            ws.cell(row=row, column=1, value="Actionnariat non détaillé dans les documents analysés.").font = F_NOTE
            row += 2
        # dirigeants
        self._header(ws, row, ["Dirigeant", "Fonction"])
        row += 1
        for p in k.managers or []:
            self._cell(ws, row, 1, p.name)
            self._cell(ws, row, 2, p.role)
            row += 1
        if not k.managers:
            ws.cell(row=row, column=1, value="Non lu dans les documents.").font = F_NOTE
            row += 1
        row += 1
        self._header(ws, row, ["Filiale / participation", "Forme", "SIREN", "% de détention", "Commune", "Pays"])
        row += 1
        for s in k.subsidiaries:
            for c_, val in enumerate([s.name, s.form, s.siren, (s.pct or 0) / 100 if s.pct is not None else None, s.city, s.country], 1):
                self._cell(ws, row, c_, val, PCT if c_ == 4 else None)
            row += 1
        if not k.subsidiaries:
            ws.cell(row=row, column=1, value="Aucune filiale détaillée dans les documents.").font = F_NOTE
        self._widths(ws, {"A": 40, "B": 16, "C": 14, "D": 18, "E": 16, "F": 14, "G": 18, "H": 38})
        self.sh_rows = (first, last) if assoc else None
        self.sh_sheet = ws

    def shareholder_charts(self):
        ws = self.wb.create_sheet("Graphiques actionnaires")
        self._title(ws, f"{self.v.name} — Graphiques de l'actionnariat",
                    "Construits à partir de l'onglet Actionnaires (mis à jour si on y corrige les parts ou dividendes).")
        if not self.sh_rows:
            ws["A4"] = "Pas de répartition du capital disponible dans les documents."
            ws["A4"].font = F_NOTE
            return
        first, last = self.sh_rows
        src = self.sh_sheet
        pie = PieChart()
        pie.title = "Répartition du capital"
        pie.add_data(Reference(src, min_col=3, min_row=first - 1, max_row=last), titles_from_data=True)
        pie.set_categories(Reference(src, min_col=1, min_row=first, max_row=last))
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True
        pie.dataLabels.showVal = False
        pie.dataLabels.showCatName = False
        pie.dataLabels.showSerName = False
        pie.dataLabels.showLegendKey = False
        pie.height, pie.width = 9, 13
        ws.add_chart(pie, "A4")
        if any(d.dividend for d in self.v.kyc.dividends_by_person):
            bar = BarChart()
            bar.type = "bar"
            bar.title = "Part du capital et part des dividendes"
            bar.add_data(Reference(src, min_col=3, min_row=first - 1, max_row=last), titles_from_data=True)
            bar.add_data(Reference(src, min_col=5, min_row=first - 1, max_row=last), titles_from_data=True)
            bar.set_categories(Reference(src, min_col=1, min_row=first, max_row=last))
            bar.y_axis.number_format = "0%"
            bar.y_axis.title = "%"
            bar.height, bar.width = 9, 16
            ws.add_chart(bar, "A24")

    # ------------------------------------------------------------------ synthèse
    def synthese(self, ws):
        v, res, f = self.v, self.res, self.f
        k = v.kyc
        ws.title = "Synthèse"
        self._title(ws, v.name, f"{FORMAT_NAMES.get(v.ex.fmt, v.ex.fmt)} — exercices : {', '.join(self.periods)}")
        from .summary import describe_activity, geography, vigilance_level
        geo = geography(v)
        vig, vig_why = vigilance_level(res, v)
        info = [("SIREN", k.siren), ("Forme juridique", k.legal_form), ("Siège", k.address), ("Code APE", k.ape),
                ("Activité principale", describe_activity(v, res)),
                ("Pays d'opération", " | ".join(f"{g_['pays']} ({g_['rôle']}{', ' + g_['risque'] if g_['risque'] != 'standard' else ''})"
                                                  for g_ in geo) or "non déterminés"),
                ("Vigilance LCB-FT suggérée", f"{vig} — {vig_why}"), ("Clôture du dernier exercice", k.closing),
                ("Opinion d'audit", k.audit_opinion), ("Statut réglementé", ", ".join(k.regulated)),
                ("Santé financière (0-100)", f"{res['health']['score']} — {res['health']['label']}"),
                ("Risque LCB-FT indicatif (0-100)", f"{res['aml']['score']} — {res['aml']['level']}")]
        r_ = 4
        for a, b in info:
            self._cell(ws, r_, 1, a, font=F_BOLD)
            self._cell(ws, r_, 2, b or "—")
            r_ += 1
        r_ += 1
        self._header(ws, r_, ["Chiffres clés"] + self.periods)
        r_ += 1
        key_rows = {}
        keys = ["pnb", "net_income", "equity", "total_assets"] if v.ex.fmt == "BANK" else \
            ["revenue", "ebit", "net_income", "equity", "total_assets", "liquid", "financial_debt_total"]
        for key in keys:
            if key not in self.refs:
                continue
            self._cell(ws, r_, 1, label(key))
            for j, p in enumerate(self.periods):
                ref = self.refs[key].get(p)
                self._cell(ws, r_, 2 + j, f"={ref}" if ref else None, NUM)
            key_rows[key] = r_
            r_ += 1
        # graphique CA / résultat net
        rev_key = "pnb" if v.ex.fmt == "BANK" else "revenue"
        if rev_key in key_rows and len(self.periods) >= 1:
            bar = BarChart()
            bar.title = "Évolution du chiffre d'affaires et du résultat net"
            rr = key_rows[rev_key]
            bar.add_data(Reference(ws, min_col=1, max_col=1 + len(self.periods), min_row=rr), from_rows=True, titles_from_data=True)
            bar.set_categories(Reference(ws, min_col=2, max_col=1 + len(self.periods), min_row=r_ - len(key_rows) - 1))
            bar.y_axis.title = v.ex.currency
            bar.y_axis.number_format = "#,##0"
            bar.y_axis.scaling.min = 0
            bar.y_axis.delete = False
            bar.x_axis.delete = False
            if "net_income" in key_rows:
                line = LineChart()
                line.add_data(Reference(ws, min_col=1, max_col=1 + len(self.periods), min_row=key_rows["net_income"]),
                              from_rows=True, titles_from_data=True)
                line.y_axis.axId = 200
                line.y_axis.title = "Résultat net"
                line.y_axis.crosses = "max"
                bar += line
            bar.height, bar.width = 9, 20
            ws.add_chart(bar, f"A{r_ + 2}")
            r_ += 21
        r_ += 1
        self._cell(ws, r_, 1, "Points clés", font=F_BOLD)
        r_ += 1
        for x in res["health"]["insights"]:
            self._cell(ws, r_, 1, "• " + x)
            ws.merge_cells(start_row=r_, start_column=1, end_row=r_, end_column=6)
            r_ += 1
        r_ += 1
        self._cell(ws, r_, 1, "Documents utilisés", font=F_BOLD)
        r_ += 1
        for s in v.sources:
            self._cell(ws, r_, 1, s["fichier"])
            self._cell(ws, r_, 2, s["périodes"])
            r_ += 1
        r_ += 1
        ws.cell(row=r_, column=1, value="Analyse automatique à partir des documents fournis, après validation de l'analyste. "
                "Les signaux LCB-FT orientent les diligences et ne constituent pas une qualification de soupçon.").font = F_NOTE
        self._widths(ws, {"A": 34, "B": 22, **{get_column_letter(3 + j): 18 for j in range(len(self.periods))}})
        for row in ws.iter_rows(min_row=4, max_row=16, min_col=2, max_col=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions["B"].width = 60

    def aml(self):
        ws = self.wb.create_sheet("LCB-FT")
        a = self.res["aml"]
        self._title(ws, f"{self.v.name} — Lecture LCB-FT", f"Score indicatif {a['score']} / 100 — {a['level']}. "
                    f"Listes GAFI mises à jour le {a.get('lists_date')}.")
        self._header(ws, 4, ["Niveau", "Catégorie", "Signal", "Constat", "Pourquoi c'est pertinent", "Diligence suggérée"])
        colors = {"élevé": "B42318", "modéré": "A55B06", "faible": "8A6D00", "atténuant": "11765F", "info": "5A6475"}
        for i, s in enumerate(a["signals"]):
            r_ = 5 + i
            for c_, val in enumerate([s["niveau"], s["catégorie"], s["signal"], s["constat"], s["pourquoi"], s["diligence"]], 1):
                c = self._cell(ws, r_, c_, val)
                c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=r_, column=1).font = Font(name="Arial", size=10, bold=True, color=colors.get(s["niveau"], "000000"))
        from .summary import geography
        base = 6 + len(a["signals"])
        geo = geography(self.v)
        self._header(ws, base, ["Pays / zone", "Rôle", "Classement (GAFI, sanctions, offshore)"])
        for i, g_ in enumerate(geo):
            self._cell(ws, base + 1 + i, 1, g_["pays"])
            self._cell(ws, base + 1 + i, 2, g_["rôle"])
            c = self._cell(ws, base + 1 + i, 3, g_["risque"])
            if g_["risque"] != "standard":
                c.font = Font(name="Arial", size=10, bold=True, color="B42318")
        base += len(geo) + 2
        self._header(ws, base, ["Noms à filtrer", "Qualité"])
        for i, row in enumerate(a["screening"]):
            self._cell(ws, base + 1 + i, 1, row["nom"])
            self._cell(ws, base + 1 + i, 2, row["qualité"])
        self._widths(ws, {"A": 12, "B": 22, "C": 36, "D": 50, "E": 50, "F": 50})

    # ------------------------------------------------------------------ build
    def build(self, docs: dict) -> bytes:
        first = self.wb.active
        self.statement("Bilan", [("Actif", BILAN_ACTIF + ["liquid"]), ("Passif", BILAN_PASSIF + ["financial_debt_total"]),
                                 ("Fonds de la clientèle (annexe)", ANNEXE)],
                       "Bilan sur tous les exercices disponibles (le document le plus récent prime pour une année présente deux fois).")
        self.statement("Compte de résultat", [("Compte de résultat", RESULTAT + ["ebitda", "general_expenses"])],
                       "Charges en valeur positive. EBITDA = résultat d'exploitation + dotations (formule).")
        self.flux()
        self.ratios()
        self.documents(docs)
        self.shareholders()
        self.shareholder_charts()
        self.aml()
        self.synthese(first)
        self.wb.move_sheet("Synthèse", offset=-self.wb.index(self.wb["Synthèse"]))
        for ws in self.wb.worksheets:
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.sheet_properties.tabColor = NAVY if ws.title in ("Synthèse", "Bilan", "Compte de résultat", "Flux de trésorerie") else "9FB3CC"
        out = io.BytesIO()
        self.wb.save(out)
        return out.getvalue()


def build_xlsx(view, res: dict, docs: dict) -> bytes:
    return Book(view, res).build(docs)
