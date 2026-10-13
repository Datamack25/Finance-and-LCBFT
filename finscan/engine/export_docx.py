"""Note de synthèse Word d'une page : comprendre l'entreprise en un coup d'œil."""
from __future__ import annotations

import io

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

from .fields import FORMAT_NAMES  # noqa: E402
from .fmt import growth, pct, symbol  # noqa: E402
from .kyc import APE_RISK  # noqa: E402
from .report import label  # noqa: E402

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
MUTED = RGBColor(0x5A, 0x64, 0x75)
GOOD = RGBColor(0x11, 0x76, 0x5F)
WARN = RGBColor(0xA5, 0x5B, 0x06)
BAD = RGBColor(0xB4, 0x23, 0x18)
FONT = "Arial"


def _short(v: float | None, sym: str) -> str:
    if v is None:
        return "—"
    a = abs(v)
    if a >= 1e9:
        s = f"{v / 1e9:.1f} Md"
    elif a >= 1e6:
        s = f"{v / 1e6:.1f} M"
    elif a >= 1e4:
        s = f"{v / 1e3:.0f} k"
    else:
        s = f"{v:,.0f} ".replace(",", " ")
    return (s.replace(".", ",") + sym).replace("-", "−")


def _shade(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _cell_margins(table, top=40, bottom=40, left=80, right=80) -> None:
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def _no_borders(table) -> None:
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for b in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{b}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def _bottom_rule(paragraph, color="1F3A5F", size=8) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(size))
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), color)
    bdr.append(b)
    pPr.append(bdr)


def _para(doc_or_cell, text="", size=9, bold=False, color=None, space_after=2, align=None, italic=False):
    p = doc_or_cell.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    if align:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.font.size, r.font.bold, r.font.italic = Pt(size), bold, italic
        if color:
            r.font.color.rgb = color
    return p


def _run(p, text, size=9, bold=False, color=None, italic=False):
    r = p.add_run(text)
    r.font.size, r.font.bold, r.font.italic = Pt(size), bold, italic
    if color:
        r.font.color.rgb = color
    return r


def _heading(doc, text):
    p = _para(doc, text, size=10.5, bold=True, color=NAVY, space_after=1)
    p.paragraph_format.space_before = Pt(5)
    _bottom_rule(p, "D9DEE6", 4)
    return p


def _chart_png(f, ex, size=(7.2, 2.7)) -> bytes | None:
    rev_key = "pnb" if ex.fmt == "BANK" else "revenue"
    periods = list(reversed(f.periods))
    rev = [f.get(rev_key, p) for p in periods]
    ni = [f.get("net_income", p) for p in periods]
    if not any(v is not None for v in rev + ni):
        return None
    sym = symbol(ex.currency)
    fig, ax = plt.subplots(figsize=size, dpi=220)
    xs = list(range(len(periods)))
    bars = ax.bar(xs, [v or 0 for v in rev], width=0.5, color="#9fb3cc", label="Chiffre d'affaires" if ex.fmt != "BANK" else "PNB")
    top = max([v for v in rev if v is not None] or [1])
    for i, (b, v) in enumerate(zip(bars, rev)):
        if v is None:
            continue
        g_ = growth(v, rev[i - 1]) if i and f.months.get(periods[i], 12) == f.months.get(periods[i - 1], 12) else None
        txt = _short(v, sym) + (f"\n{pct(g_)}" if g_ is not None else "")
        ax.text(b.get_x() + b.get_width() / 2, v + top * 0.03, txt, ha="center", va="bottom", fontsize=7, color="#1f3a5f",
                fontweight="bold")
    ax.set_ylim(min(0, min([v for v in rev if v is not None] or [0])) * 1.2, top * 1.45)
    ax.set_xticks(xs, periods, fontsize=8)
    ax.tick_params(axis="y", labelsize=7, colors="#5a6475")
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda x, _: _short(x, "")))
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.spines["left"].set_color("#d9dee6")
    ax.spines["bottom"].set_color("#d9dee6")
    if any(v is not None for v in ni):
        ax2 = ax.twinx()
        pts = [(x, v) for x, v in zip(xs, ni) if v is not None]
        ax2.plot([x for x, _ in pts], [v for _, v in pts], color="#11765f", marker="o", lw=2, label="Résultat net")
        for x, v in pts:
            ax2.annotate(_short(v, sym), (x, v), textcoords="offset points", xytext=(0, -12), ha="center", fontsize=7,
                         color="#b42318" if v < 0 else "#11765f")
        lo, hi = min(0, min(v for _, v in pts)), max(0, max(v for _, v in pts))
        span = (hi - lo) or 1
        ax2.set_ylim(lo - span * 0.25 if lo < 0 else 0, hi + span * 1.0)
        ax2.tick_params(axis="y", labelsize=7, colors="#11765f")
        ax2.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda x, _: _short(x, "")))
        for s in ("top", "left"):
            ax2.spines[s].set_visible(False)
        ax2.spines["right"].set_color("#d9dee6")
        h1, l1 = ax.get_legend_handles_labels()
        h2, l2 = ax2.get_legend_handles_labels()
        ax.legend(h1 + h2, l1 + l2, loc="upper left", fontsize=7, frameon=False, ncol=2)
    fig.tight_layout(pad=0.4)
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return buf.getvalue()


def _clip(s: str, n: int) -> str:
    s = " ".join((s or "").split())
    return s if len(s) <= n else s[: n - 1].rsplit(" ", 1)[0] + "…"


def _cell_text(cell, first=True):
    """Premier paragraphe vide de la cellule, réutilisé pour ne pas laisser de ligne blanche."""
    p = cell.paragraphs[0]
    p.text = ""
    p.paragraph_format.space_after = Pt(1)
    return p


def _col_heading(cell, text, color_hex, bg_hex):
    p = _cell_text(cell)
    p.paragraph_format.space_after = Pt(3)
    _run(p, text, 10.5, True, RGBColor.from_string(color_hex))
    _bottom_rule(p, color_hex, 8)


def _sub(cell, text):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    _run(p, text, 8.5, True, NAVY)
    return p


def _line(cell, text, size=9, color=None, bold=False, bullet=False):
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    if bullet:
        p.paragraph_format.left_indent = Cm(0.3)
        p.paragraph_format.first_line_indent = Cm(-0.3)
        _run(p, "• ", size, True, NAVY)
    _run(p, text, size, bold, color)
    return p


def build_docx(view, res: dict, author: str = "") -> bytes:
    """Note d'une page, équilibrée : lecture financière (gauche) / lecture LCB-FT (droite)."""
    import datetime as dt

    from .fmt import ratio as fmt_ratio
    from .kyc import is_legal_entity
    from .summary import describe_activity, geography, vigilance_level

    f, h, a = res["fin"], res["health"], res["aml"]
    ex, k = view.ex, view.kyc
    sym = symbol(ex.currency)
    periods = list(reversed(f.periods))
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.4)
    sec.top_margin, sec.bottom_margin = Cm(1.2), Cm(1.0)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(8.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    # ---------------------------------------------------------------- en-tête
    _para(doc, view.name, size=16, bold=True, color=NAVY, space_after=0)
    meta = [x for x in (k.legal_form, f"SIREN {k.siren}" if k.siren else "", k.address, f"APE {k.ape}" if k.ape else "",
                        f"Exercices {periods[0]}–{periods[-1]}" if len(periods) > 1 else (f"Exercice {periods[0]}" if periods else ""))
            if x]
    p = _para(doc, "   |   ".join(meta), size=8, color=MUTED, space_after=3)
    _bottom_rule(p)

    # ---------------------------------------------------------------- scores
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _no_borders(t)
    _cell_margins(t, 50, 50, 110, 110)
    hs, hl = h["score"], h["label"] or "données insuffisantes"
    hc = GOOD if (hs or 0) >= 70 else WARN if (hs or 0) >= 50 else BAD
    ac = GOOD if a["level"] == "Faible" else WARN if a["level"] == "Modéré" else BAD
    vig, vig_why = vigilance_level(res, view)
    for cell, title, score, lab, col, bg, extra in (
            (t.cell(0, 0), "Santé financière", hs, hl, hc, "EEF4F1", ""),
            (t.cell(0, 1), "Risque LCB-FT indicatif", a["score"], a["level"], ac, "F7EFEA", f"   vigilance {vig}")):
        _shade(cell, bg)
        pp = _cell_text(cell)
        _run(pp, title + "   ", 8.5, color=MUTED)
        _run(pp, f"{'—' if score is None else int(score)}", 14, True, col)
        _run(pp, " / 100   ", 8.5, color=MUTED)
        _run(pp, lab, 10, True, col)
        if extra:
            _run(pp, extra, 8, True, col)

    # ---------------------------------------------------------------- profil (commun)
    p = _para(doc, space_after=1)
    p.paragraph_format.space_before = Pt(5)
    _run(p, "Activité  ", 8.5, True, NAVY)
    _run(p, _clip(describe_activity(view, res), 420), 8.5)
    geo = geography(view)
    p = _para(doc, space_after=3)
    _run(p, "Pays d'opération  ", 8.5, True, NAVY)
    if geo:
        for i, g_ in enumerate(geo[:8]):
            if i:
                _run(p, "  |  ", 8.5, color=MUTED)
            _run(p, g_["pays"], 8.5, True)
            _run(p, f" ({g_['rôle']})", 8, color=MUTED)
            if g_["risque"] != "standard":
                _run(p, f" {g_['risque']}", 8, True, BAD)
        risky = [g_ for g_ in geo if g_["risque"] != "standard"]
        _run(p, "   —   " + ("pays à risque identifiés" if risky else "aucun pays sur les listes GAFI, sanctions ou offshore"),
             8, not risky and False, BAD if risky else GOOD)
    else:
        _run(p, "non déterminés à partir des documents", 8.5, color=MUTED)

    # ---------------------------------------------------------------- deux colonnes
    cols = doc.add_table(rows=1, cols=2)
    _no_borders(cols)
    _cell_margins(cols, 20, 20, 60, 60)
    left, right = cols.cell(0, 0), cols.cell(0, 1)
    for c in (left, right):
        c.width = Cm(9.1)

    # ----- gauche : lecture financière
    _col_heading(left, "Lecture financière", "1F3A5F", "FFFFFF")
    keys = ["pnb", "net_income", "equity", "client_funds"] if ex.fmt == "BANK" else \
        ["revenue", "ebit", "net_income", "equity", "liquid", "financial_debt_total"]
    keys = [kk for kk in keys if any(f.get(kk, pp) is not None for pp in periods)][:5]
    shown = periods[-3:]
    mt = left.add_table(rows=1 + len(keys), cols=1 + len(shown))
    mt.style = "Table Grid"
    _cell_margins(mt, 15, 15, 50, 50)
    for j, txt in enumerate(["en " + sym] + shown):
        c = mt.cell(0, j)
        _shade(c, "1F3A5F")
        pp = _cell_text(c)
        pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT if j else WD_ALIGN_PARAGRAPH.LEFT
        _run(pp, txt, 7.5, True, RGBColor(0xFF, 0xFF, 0xFF))
    for i, kk in enumerate(keys, 1):
        pp = _cell_text(mt.cell(i, 0))
        _run(pp, label(kk).replace(" (opérations avec la clientèle)", ""), 7.5, kk in ("revenue", "pnb", "net_income"))
        for j, pr in enumerate(shown, 1):
            v = f.get(kk, pr)
            q = f.prev(pr)
            g_ = growth(v, f.get(kk, q)) if q else None
            pp = _cell_text(mt.cell(i, j))
            pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _run(pp, _short(v, ""), 7.5)
            if g_ is not None:
                inv = kk == "financial_debt_total"
                colr = (GOOD if (g_ < 0 if inv else g_ > 0) else BAD) if abs(g_) > 0.005 else MUTED
                _run(pp, f" {pct(g_)}", 6.5, True, colr)
    widths = [Cm(3.0)] + [Cm(5.0 / len(shown))] * len(shown)
    for row in mt.rows:
        for c, w in zip(row.cells, widths):
            c.width = w
    png = _chart_png(f, ex, size=(4.4, 3.1))
    if png:
        pp = left.add_paragraph()
        pp.paragraph_format.space_before = Pt(3)
        pp.paragraph_format.space_after = Pt(0)
        pp.add_run().add_picture(io.BytesIO(png), width=Cm(8.8))
    r = {x["ratio"]: x for x in res["ratios"]}
    pick = ([("Coefficient d'exploitation", "Coef. d'exploitation"), ("ROE", "ROE"), ("Capitaux propres / total bilan", "Fonds propres / bilan"),
             ("Comptes dédiés (cantonnement + acquisition + réserves) / fonds clientèle", "Couverture fonds clients")]
            if ex.fmt == "BANK" else
            [("Marge nette", "Marge nette"), ("ROE (rentabilité des capitaux propres)", "ROE"),
             ("Autonomie financière (capitaux propres / total bilan)", "Autonomie fin."),
             ("Gearing (dettes financières / capitaux propres)", "Gearing"),
             ("Liquidité générale (actif circulant / passif courant)", "Liquidité gén.")])
    vals = [f"{short} {fmt_ratio(r[name]['N'], r[name]['format'])}" for name, short in pick if name in r and r[name]["N"] is not None]
    if h.get("altman"):
        vals.append(f"Z-score {h['altman']['z']:.2f}".replace(".", ","))
    if vals:
        _sub(left, "Ratios du dernier exercice")
        _line(left, "   |   ".join(vals), 8, NAVY, True)
    _sub(left, "Points clés")
    for ins in h["insights"][:5]:
        _line(left, _clip(ins, 120), 9, bullet=True)

    # ----- droite : lecture LCB-FT
    _col_heading(right, "Lecture LCB-FT", "8C2A1C", "FFFFFF")
    pp = right.add_paragraph()
    pp.paragraph_format.space_after = Pt(2)
    _run(pp, f"Vigilance suggérée : {vig}. ", 8.5, True, ac)
    _run(pp, _clip(vig_why, 150), 8.5)

    _sub(right, "Actionnariat et bénéficiaires effectifs")
    if k.associates:
        for p_ in k.associates[:5]:
            legal = is_legal_entity(p_.name)
            pct_txt = f"{p_.pct:g} %" if p_.pct is not None else "% n/d"
            if legal:
                note = " — personne morale : remonter la chaîne de détention jusqu'aux personnes physiques"
            elif (p_.pct or 0) > 25:
                note = " — bénéficiaire effectif présumé"
            else:
                note = ""
            _line(right, f"{p_.name} {pct_txt}{note}", 8.5, bullet=True)
    else:
        _line(right, "Actionnariat non détaillé dans les documents : à obtenir (statuts, RBE).", 8.5, MUTED, bullet=True)
    if k.managers:
        _line(right, "Dirigeant(s) : " + ", ".join(f"{p_.name} ({p_.role.lower()})" for p_ in k.managers[:3]), 8.5)

    # profil de risque par catégorie (niveau le plus élevé rencontré)
    order = {"élevé": 0, "modéré": 1, "faible": 2, "info": 3, "atténuant": 4}
    cats = {}
    for s_ in a["signals"]:
        c_ = s_["catégorie"]
        if c_ not in cats or order[s_["niveau"]] < order[cats[c_]]:
            cats[c_] = s_["niveau"]
    for c_ in ("Actionnariat & bénéficiaires effectifs", "Structure & activité", "Flux & trésorerie", "Cohérence économique",
               "Géographie", "Secteur", "Gouvernance & conformité"):
        cats.setdefault(c_, "aucun signal")
    _sub(right, "Profil de risque par axe")
    rt = right.add_table(rows=4, cols=4)
    _no_borders(rt)
    _cell_margins(rt, 10, 10, 30, 30)
    items = list(cats.items())[:8]
    short_c = {"Actionnariat & bénéficiaires effectifs": "Actionnariat / BE", "Structure & activité": "Structure",
               "Flux & trésorerie": "Flux", "Cohérence économique": "Cohérence éco.", "Gouvernance & conformité": "Gouvernance"}
    lvl_col = {"élevé": BAD, "modéré": WARN, "faible": RGBColor(0x8A, 0x6D, 0x00), "atténuant": GOOD, "info": MUTED, "aucun signal": GOOD}
    for i, (c_, lv) in enumerate(items):
        row, col = i % 4, (i // 4) * 2
        pp = _cell_text(rt.cell(row, col))
        _run(pp, short_c.get(c_, c_), 8)
        pp = _cell_text(rt.cell(row, col + 1))
        _run(pp, lv, 8, True, lvl_col.get(lv, MUTED))
    for row in rt.rows:
        for c, w in zip(row.cells, [Cm(2.6), Cm(1.9), Cm(2.6), Cm(1.9)]):
            c.width = w

    _sub(right, "Signaux d'attention")
    tops = [s for s in a["signals"] if s["niveau"] in ("élevé", "modéré")][:5]
    if len(tops) < 3:
        tops += [s for s in a["signals"] if s["niveau"] in ("faible", "atténuant")][:3 - len(tops)]
    colors = {"élevé": BAD, "modéré": WARN, "faible": MUTED, "atténuant": GOOD, "info": MUTED}
    for s in tops:
        pp = right.add_paragraph()
        pp.paragraph_format.space_after = Pt(1)
        pp.paragraph_format.line_spacing = 1.0
        _run(pp, s["niveau"].capitalize() + "  ", 8, True, colors[s["niveau"]])
        _run(pp, s["signal"] + " : ", 9, True)
        _run(pp, _clip(s["constat"], 115), 9)

    dil = []
    for s in [x for x in a["signals"] if x["niveau"] in ("élevé", "modéré", "faible")]:
        d = _clip(s["diligence"].split(" ; ")[0].split(". ")[0], 105)
        if d not in dil:
            dil.append(d)
    if dil:
        _sub(right, "Diligences prioritaires")
        for d in dil[:4]:
            _line(right, d, 9, bullet=True)
    status = []
    if k.regulated:
        status.append(k.regulated[0])
    if k.listed:
        status.append("société cotée")
    status.append(f"comptes : {k.audit_opinion}")
    _line(right, "Contrôle et statut : " + ", ".join(status) + f". {len(a['screening'])} nom(s) à filtrer (gel des avoirs, sanctions, PPE).",
          8, MUTED)

    # ---------------------------------------------------------------- conclusion
    p = _para(doc, space_after=1)
    p.paragraph_format.space_before = Pt(6)
    _bottom_rule(p, "D9DEE6", 4)
    _run(p, "Conclusion", 10, True, NAVY)
    rev_g = growth(f.get("pnb" if ex.fmt == "BANK" else "revenue"), f.get("pnb" if ex.fmt == "BANK" else "revenue", "n1"))
    fin_txt = (f"Situation financière {hl.lower()} ({hs}/100)" if hs is not None else "Situation financière non évaluée") + \
        (f", activité {'en croissance' if rev_g > 0 else 'en recul'} de {pct(rev_g).lstrip('+')} sur le dernier exercice" if rev_g is not None else "") + "."
    p = _para(doc, space_after=0)
    _run(p, "Financière : ", 8.5, True, NAVY)
    _run(p, fin_txt, 8.5)
    high = [s["signal"].lower() for s in a["signals"] if s["niveau"] == "élevé"][:2]
    aml_txt = f"risque {a['level'].lower()} ({a['score']}/100), vigilance {vig}" + \
        (f" ; points bloquants à lever : {', '.join(high)}." if high else ".")
    p = _para(doc, space_after=0)
    _run(p, "LCB-FT : ", 8.5, True, RGBColor(0x8C, 0x2A, 0x1C))
    _run(p, aml_txt, 8.5)

    # python-docx ajoute un paragraphe vide après chaque tableau imbriqué : on le retire
    for cell in (left, right):
        for tbl in cell._tc.findall(qn("w:tbl")):
            nxt = tbl.getnext()
            if nxt is not None and nxt.tag == qn("w:p") and not "".join(nxt.itertext()).strip() and not nxt.findall(".//" + qn("w:drawing")):
                cell._tc.remove(nxt)
    foot = (f"Sources : {_clip(view.file_name, 110)}. {FORMAT_NAMES.get(ex.fmt, ex.fmt)}. Note générée le "
            f"{dt.date.today().strftime('%d/%m/%Y')}{' par ' + author if author else ''} (FinScan LCB-FT). Signaux et niveau de vigilance "
            "indicatifs, à confirmer par la classification des risques de l'établissement ; ni qualification de soupçon, ni conseil en investissement.")
    p = _para(doc, foot, 6.5, color=MUTED, italic=True, space_after=0)
    p.paragraph_format.space_before = Pt(4)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
